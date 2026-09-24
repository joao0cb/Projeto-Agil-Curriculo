#!/usr/bin/env python3
"""Gera UX_NIELSEN_REPORT.docx e UX_NIELSEN_REPORT.pdf a partir do .md.

O conteúdo textual é o mesmo do relatório enviado ao usuário — o script
apenas converte Markdown (títulos, tabelas, negrito/itálico/código) para
DOCX (python-docx) e PDF (reportlab), sem alterar o texto.
"""
import re
from pathlib import Path

import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable,
)

ROOT = Path(__file__).resolve().parent.parent
MD_PATH = ROOT / "docs" / "UX_NIELSEN_REPORT.md"
DOCX_PATH = ROOT / "docs" / "UX_NIELSEN_REPORT.docx"
PDF_PATH = ROOT / "docs" / "UX_NIELSEN_REPORT.pdf"


def strip_inline(text: str) -> str:
    """Remove marcação inline (**, *, `) mantendo o texto."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text


def parse_markdown(path: Path):
    """Divide o .md em blocos: heading, paragraph, hr, table."""
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            blocks.append(("h", level, line[level:].strip()))
            i += 1
        elif line.strip() == "---":
            blocks.append(("hr",))
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
        elif line.strip():
            blocks.append(("p", line.strip()))
            i += 1
        else:
            i += 1
    return blocks


# ────────────────────────────── DOCX ──────────────────────────────

def add_runs(paragraph, text: str):
    """Adiciona runs preservando negrito/itálico/código inline."""
    pos = 0
    pattern = re.compile(r"(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)")
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        if match.group(2) is not None:
            run = paragraph.add_run(match.group(2))
            run.bold = True
        elif match.group(3) is not None:
            run = paragraph.add_run(match.group(3))
            run.italic = True
        else:
            run = paragraph.add_run(match.group(4))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def build_docx(blocks):
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    for block in blocks:
        kind = block[0]
        if kind == "h":
            level, text = block[1], block[2]
            heading = document.add_heading(strip_inline(text), level=min(level, 3))
            for run in heading.runs:
                run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        elif kind == "hr":
            document.add_paragraph("─" * 60).alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "table":
            rows = block[1]
            ncols = max(len(r) for r in rows)
            table = document.add_table(rows=len(rows), cols=ncols)
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c in range(ncols):
                    cell = table.cell(r, c)
                    cell.text = ""
                    para = cell.paragraphs[0]
                    add_runs(para, row[c] if c < len(row) else "")
                    for run in para.runs:
                        run.font.size = Pt(9)
                        if r == 0:
                            run.bold = True
            document.add_paragraph()
        else:
            para = document.add_paragraph()
            add_runs(para, block[1])

    document.save(DOCX_PATH)


# ─────────────────────────────── PDF ───────────────────────────────

def esc(text: str) -> str:
    """Escapa XML e converte marcação inline para tags do Paragraph."""
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
    text = re.sub(r"`(.+?)`", r"<font face='Courier' size='8'>\1</font>", text)
    return text


def build_pdf(blocks):
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=15, spaceAfter=8, textColor=colors.HexColor("#111827"))
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=6, textColor=colors.HexColor("#111827"))
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontSize=9, leading=12, alignment=TA_LEFT)
    cell_style = ParagraphStyle("Cell", parent=body, fontSize=7.5, leading=9.5)
    cell_head = ParagraphStyle("CellHead", parent=cell_style, fontSize=7.5, leading=9.5)

    story = []
    for block in blocks:
        kind = block[0]
        if kind == "h":
            level, text = block[1], block[2]
            story.append(Paragraph(esc(strip_inline(text)), h1 if level <= 1 else h2))
        elif kind == "hr":
            story.append(Spacer(1, 2))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#9CA3AF")))
            story.append(Spacer(1, 4))
        elif kind == "table":
            rows = block[1]
            ncols = max(len(r) for r in rows)
            data = []
            for r, row in enumerate(rows):
                padded = row + [""] * (ncols - len(row))
                st = cell_head if r == 0 else cell_style
                data.append([Paragraph(esc(strip_inline(c)), st) for c in padded])
            # Larguras: colunas de status estreitas, "onde" larga.
            if ncols == 3:
                col_widths = [46 * mm, 22 * mm, 112 * mm]
            elif ncols == 2:
                col_widths = [62 * mm, 118 * mm]
            else:
                col_widths = [180 * mm / ncols] * ncols
            table = Table(data, colWidths=col_widths, repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F4F6")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ]))
            story.append(table)
            story.append(Spacer(1, 6))
        else:
            story.append(Paragraph(esc(block[1]), body))
            story.append(Spacer(1, 3))

    doc = SimpleDocTemplate(
        str(PDF_PATH), pagesize=A4,
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=15 * mm,
        title="Relatório — 10 Heurísticas de Nielsen por Tela",
    )
    doc.build(story)


if __name__ == "__main__":
    blocks = parse_markdown(MD_PATH)
    build_docx(blocks)
    build_pdf(blocks)
    print(f"OK: {DOCX_PATH.name}, {PDF_PATH.name}")
